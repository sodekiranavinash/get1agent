from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}


@dataclass
class ExtractedImage:
    data: bytes
    extension: str
    page: int | None = None
    width: int | None = None
    height: int | None = None


@dataclass
class Extraction:
    text: str
    images: list[ExtractedImage] = field(default_factory=list)
    # Shape metadata for the ingestion timeline. Only the fields that make sense
    # for a given format are populated (e.g. pages for PDF, rows/sheets for
    # spreadsheets, paragraphs for DOCX); the rest stay None and are omitted.
    pages: int | None = None
    rows: int | None = None
    sheets: int | None = None
    paragraphs: int | None = None
    # Per-page text for paginated formats (PDF). ``None`` when the format has no
    # page concept; the chunker uses it to tag chunks with their source page.
    page_texts: list[str] | None = None


class UnsupportedDocument(Exception):
    pass


def extension_of(file_name: str) -> str:
    return os.path.splitext(file_name or "")[1].lower()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _table_to_markdown(rows: list[list[str]]) -> str:
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> Extraction:
    try:
        import pymupdf as fitz  # type: ignore[import-not-found]
    except ImportError:
        try:
            import fitz  # type: ignore[import-not-found,no-redef]
        except ImportError as exc:  # pragma: no cover - depends on packaging
            raise UnsupportedDocument("PDF extraction requires pymupdf") from exc

    document = fitz.open(stream=data, filetype="pdf")
    try:
        pages: list[str] = []
        images: list[ExtractedImage] = []
        for index, page in enumerate(document):
            pages.append(page.get_text("text"))
            for info in page.get_images(full=True):
                try:
                    raw = document.extract_image(info[0])
                except Exception:  # noqa: BLE001 - skip unreadable image
                    continue
                images.append(
                    ExtractedImage(
                        data=raw["image"],
                        extension=raw.get("ext", "png"),
                        page=index + 1,
                        width=raw.get("width"),
                        height=raw.get("height"),
                    )
                )
        return Extraction(
            "\n\n".join(pages).strip(),
            images,
            pages=len(pages),
            page_texts=pages,
        )
    finally:
        document.close()


def _extract_docx(data: bytes) -> Extraction:
    try:
        from docx import Document as DocxDocument  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - depends on packaging
        raise UnsupportedDocument("DOCX extraction requires python-docx") from exc

    document = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    blocks: list[str] = list(paragraphs)
    table_rows = 0
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        table_rows += len(rows)
        markdown = _table_to_markdown(rows)
        if markdown:
            blocks.append(markdown)

    images: list[ExtractedImage] = []
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        images.append(
            ExtractedImage(data=part.blob, extension=str(part.partname.ext or "png"))
        )

    return Extraction(
        "\n\n".join(blocks).strip(),
        images,
        paragraphs=len(paragraphs) or None,
        rows=table_rows or None,
    )


def _extract_xlsx(data: bytes) -> Extraction:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - depends on packaging
        raise UnsupportedDocument("XLSX extraction requires openpyxl") from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        sheet_count = 0
        row_count = 0
        for sheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            row_count += len(rows)
            markdown = _table_to_markdown(rows)
            if markdown:
                sheets.append(f"## {sheet.title}\n\n{markdown}")
                sheet_count += 1
        return Extraction(
            "\n\n".join(sheets).strip(),
            sheets=sheet_count or None,
            rows=row_count or None,
        )
    finally:
        workbook.close()


def _extract_csv(data: bytes) -> Extraction:
    reader = csv.reader(io.StringIO(_decode(data)))
    rows = [[cell.strip() for cell in row] for row in reader]
    return Extraction(_table_to_markdown(rows), rows=len(rows) or None)


def extract(data: bytes, file_name: str) -> Extraction:
    """Extract normalized markdown text and embedded images from a document."""
    extension = extension_of(file_name)
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".xlsx":
        return _extract_xlsx(data)
    if extension == ".csv":
        return _extract_csv(data)
    if extension in TEXT_EXTENSIONS or not extension:
        return Extraction(_decode(data).strip())
    raise UnsupportedDocument(f"Unsupported file type: {extension}")
