from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass, field

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}


@dataclass
class ExtractedImage:
    data: bytes
    extension: str
    page: int | None = None
    width: int | None = None
    height: int | None = None


@dataclass
class Extraction:
    text: str
    images: list[ExtractedImage] = field(default_factory=list)


class UnsupportedDocument(Exception):
    pass


def extension_of(file_name: str) -> str:
    return os.path.splitext(file_name or "")[1].lower()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _table_to_markdown(rows: list[list[str]]) -> str:
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> Extraction:
    try:
        import pymupdf as fitz  # type: ignore[import-not-found]
    except ImportError:
        try:
            import fitz  # type: ignore[import-not-found,no-redef]
        except ImportError as exc:  # pragma: no cover - depends on packaging
            raise UnsupportedDocument("PDF extraction requires pymupdf") from exc

    document = fitz.open(stream=data, filetype="pdf")
    try:
        pages: list[str] = []
        images: list[ExtractedImage] = []
        for index, page in enumerate(document):
            pages.append(page.get_text("text"))
            for info in page.get_images(full=True):
                try:
                    raw = document.extract_image(info[0])
                except Exception:  # noqa: BLE001 - skip unreadable image
                    continue
                images.append(
                    ExtractedImage(
                        data=raw["image"],
                        extension=raw.get("ext", "png"),
                        page=index + 1,
                        width=raw.get("width"),
                        height=raw.get("height"),
                    )
                )
        return Extraction("\n\n".join(pages).strip(), images)
    finally:
        document.close()


def _extract_docx(data: bytes) -> Extraction:
    try:
        from docx import Document as DocxDocument  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - depends on packaging
        raise UnsupportedDocument("DOCX extraction requires python-docx") from exc

    document = DocxDocument(io.BytesIO(data))
    blocks: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        markdown = _table_to_markdown(rows)
        if markdown:
            blocks.append(markdown)

    images: list[ExtractedImage] = []
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        images.append(
            ExtractedImage(data=part.blob, extension=str(part.partname.ext or "png"))
        )

    return Extraction("\n\n".join(blocks).strip(), images)


def _extract_xlsx(data: bytes) -> Extraction:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - depends on packaging
        raise UnsupportedDocument("XLSX extraction requires openpyxl") from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for sheet in workbook.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            markdown = _table_to_markdown(rows)
            if markdown:
                sheets.append(f"## {sheet.title}\n\n{markdown}")
        return Extraction("\n\n".join(sheets).strip())
    finally:
        workbook.close()


def _extract_csv(data: bytes) -> Extraction:
    reader = csv.reader(io.StringIO(_decode(data)))
    rows = [[cell.strip() for cell in row] for row in reader]
    return Extraction(_table_to_markdown(rows))


def extract(data: bytes, file_name: str) -> Extraction:
    """Extract normalized markdown text and embedded images from a document."""
    extension = extension_of(file_name)
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension == ".xlsx":
        return _extract_xlsx(data)
    if extension == ".csv":
        return _extract_csv(data)
    if extension in TEXT_EXTENSIONS or not extension:
        return Extraction(_decode(data).strip())
    raise UnsupportedDocument(f"Unsupported file type: {extension}")
